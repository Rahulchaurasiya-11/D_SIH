"""End-to-end API tests: authentication, role enforcement, scanning, search, exports."""

from tests.conftest import COMPLIANT_LABEL


def scan_text(client, headers, lines, **extra):
    body = {"raw_text": "\n".join(lines), "filename": "label.txt", "source": "manual"}
    body.update(extra)
    return client.post("/api/v1/analyze-text", json=body, headers=headers)


# --------------------------------------------------------------------------- #
# Health and auth
# --------------------------------------------------------------------------- #

def test_health_is_public(client):
    response = client.get("/api/v1/health")
    assert response.status_code == 200
    assert response.json()["status"] == "online"


def test_first_account_becomes_admin(admin):
    assert admin["user"]["role"] == "ADMIN"


def test_subsequent_account_is_inspector_even_if_admin_requested(client, admin):
    """Privilege escalation via the registration payload must not be possible."""
    response = client.post("/api/v1/auth/register", json={
        "email": "sneaky@lm.gov.in", "password": "SneakyPass#2026",
        "full_name": "Escalation Attempt", "role": "ADMIN",
    })
    assert response.status_code == 201
    assert response.json()["user"]["role"] == "INSPECTOR"


def test_duplicate_email_rejected(client, inspector):
    response = client.post("/api/v1/auth/register", json={
        "email": "inspector@lm.gov.in", "password": "AnotherPass#2026", "full_name": "Clone",
    })
    assert response.status_code == 409


def test_login_succeeds_and_wrong_password_fails(client, inspector):
    ok = client.post("/api/v1/auth/login",
                     json={"email": "inspector@lm.gov.in", "password": "InspectorPass#2026"})
    assert ok.status_code == 200
    assert ok.json()["user"]["full_name"] == "Ravi Inspector"

    bad = client.post("/api/v1/auth/login",
                      json={"email": "inspector@lm.gov.in", "password": "wrong"})
    assert bad.status_code == 401


def test_unknown_user_and_wrong_password_are_indistinguishable(client):
    """The login endpoint must not reveal which officer accounts exist."""
    a = client.post("/api/v1/auth/login", json={"email": "nobody@lm.gov.in", "password": "x" * 12})
    b = client.post("/api/v1/auth/login", json={"email": "inspector@lm.gov.in", "password": "x" * 12})
    assert a.status_code == b.status_code == 401
    assert a.json()["detail"] == b.json()["detail"]


def test_refresh_token_issues_new_access_token(client, inspector):
    response = client.post("/api/v1/auth/refresh",
                           json={"refresh_token": inspector["refresh_token"]})
    assert response.status_code == 200
    assert response.json()["access_token"]


def test_access_token_rejected_as_refresh_token(client, inspector):
    """Token type confusion: an access token must not be redeemable for a refresh."""
    response = client.post("/api/v1/auth/refresh",
                           json={"refresh_token": inspector["access_token"]})
    assert response.status_code == 401


# --------------------------------------------------------------------------- #
# Access control
# --------------------------------------------------------------------------- #

def test_protected_route_without_token_is_401(client):
    assert client.get("/api/v1/inspections").status_code == 401


def test_garbage_token_is_401(client):
    response = client.get("/api/v1/inspections", headers={"Authorization": "Bearer not.a.token"})
    assert response.status_code == 401


def test_inspector_cannot_list_users(client, inspector_headers):
    assert client.get("/api/v1/auth/users", headers=inspector_headers).status_code == 403


def test_admin_can_list_users(client, admin_headers):
    response = client.get("/api/v1/auth/users", headers=admin_headers)
    assert response.status_code == 200
    assert len(response.json()) >= 2


def test_inspector_cannot_reload_rules(client, inspector_headers):
    assert client.post("/api/v1/rules/reload", headers=inspector_headers).status_code == 403


def test_admin_can_reload_rules(client, admin_headers):
    response = client.post("/api/v1/rules/reload", headers=admin_headers)
    assert response.status_code == 200
    assert response.json()["success"] is True


def test_admin_cannot_demote_self(client, admin_headers, admin):
    response = client.patch(
        "/api/v1/auth/users/%s/role?role=INSPECTOR" % admin["user"]["id"], headers=admin_headers
    )
    assert response.status_code == 400


# --------------------------------------------------------------------------- #
# Scanning
# --------------------------------------------------------------------------- #

def test_scan_text_creates_inspection(client, inspector_headers):
    response = scan_text(client, inspector_headers, COMPLIANT_LABEL)
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "COMPLIANT"
    assert body["inspection_id"].startswith("insp_")


def test_scan_rejects_empty_text(client, inspector_headers):
    response = scan_text(client, inspector_headers, [])
    assert response.status_code == 400


def test_non_compliant_scan_records_violations(client, inspector_headers):
    lines = [ln.replace(" (Inclusive of all taxes)", "") for ln in COMPLIANT_LABEL]
    body = scan_text(client, inspector_headers, lines).json()
    assert body["status"] != "COMPLIANT"
    assert body["violations"]


def test_listing_ingestion_parses_html(client, inspector_headers):
    """E-commerce listings are audited by the same engine, not a second rule set."""
    html = (
        "<div><script>var x=1;</script><h1>Almond Pack</h1>"
        "<p>Net Quantity: 500 g</p><p>MRP Rs. 450.00 (Inclusive of all taxes)</p>"
        "<p>Mfg Date: 02/2026</p><p>care@brand.com</p><p>Helpline: 1800-111-2233</p>"
        "<p>Country of Origin: India</p><span>Add to cart</span></div>"
    )
    response = client.post("/api/v1/analyze-listing", headers=inspector_headers, json={
        "listing_html": html, "source_url": "https://example.com/p/1",
        "platform": "Example", "product_title": "Almond Pack",
    })
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["extracted_metadata"]["mrp"]
    # Script contents and marketplace chrome must not reach the rule engine.
    joined = " ".join(body["raw_text_dump"])
    assert "var x=1" not in joined
    assert "Add to cart" not in joined


def test_listing_rejects_empty_payload(client, inspector_headers):
    response = client.post("/api/v1/analyze-listing", headers=inspector_headers, json={})
    assert response.status_code == 400


def test_inspector_verification_overrides_a_false_positive(client, inspector_headers):
    lines = [ln.replace(" (Inclusive of all taxes)", "") for ln in COMPLIANT_LABEL]
    created = scan_text(client, inspector_headers, lines).json()
    assert created["violations"]

    response = client.post("/api/v1/verify-and-audit", headers=inspector_headers, json={
        "inspection_id": created["inspection_id"],
        "manual_overrides": {"mrp": "450.00", "taxes_included": True},
    })
    assert response.status_code == 200, response.text
    revised = response.json()
    assert revised["is_manually_verified"] is True
    assert revised["overall_score"] >= created["overall_score"]


# --------------------------------------------------------------------------- #
# Repository: search and retrieval
# --------------------------------------------------------------------------- #

def test_search_returns_paginated_results(client, inspector_headers):
    scan_text(client, inspector_headers, COMPLIANT_LABEL)
    response = client.get("/api/v1/inspections?page=1&page_size=5", headers=inspector_headers)
    assert response.status_code == 200
    body = response.json()
    assert body["total"] >= 1
    assert len(body["items"]) <= 5
    assert body["pages"] >= 1


def test_search_filters_by_status(client, inspector_headers):
    response = client.get("/api/v1/inspections?status=COMPLIANT", headers=inspector_headers)
    assert response.status_code == 200
    assert all(item["status"] == "COMPLIANT" for item in response.json()["items"])


def test_search_filters_by_score_range(client, inspector_headers):
    response = client.get("/api/v1/inspections?min_score=90&max_score=100", headers=inspector_headers)
    assert response.status_code == 200
    assert all(90 <= item["overall_score"] <= 100 for item in response.json()["items"])


def test_inspector_sees_only_their_own_inspections(client, inspector_headers, admin_headers):
    """An inspector must not be able to read a colleague's inspections by id."""
    theirs = scan_text(client, admin_headers, COMPLIANT_LABEL).json()["inspection_id"]

    listed = client.get("/api/v1/inspections", headers=inspector_headers).json()
    assert all(item["id"] != theirs for item in listed["items"])

    direct = client.get("/api/v1/inspections/" + theirs, headers=inspector_headers)
    assert direct.status_code == 403


def test_inspector_cannot_widen_scope_via_officer_id(client, inspector_headers, admin):
    """Passing someone else's officer_id must not bypass the inspector's own-records scope."""
    response = client.get("/api/v1/inspections?officer_id=" + admin["user"]["id"],
                          headers=inspector_headers)
    assert response.status_code == 200
    assert all(item["officer_id"] != admin["user"]["id"] for item in response.json()["items"])


def test_missing_inspection_is_404(client, admin_headers):
    assert client.get("/api/v1/inspections/insp_doesnotexist", headers=admin_headers).status_code == 404


def test_dashboard_stats_shape(client, admin_headers):
    response = client.get("/api/v1/dashboard/stats?days=30", headers=admin_headers)
    assert response.status_code == 200
    stats = response.json()
    for key in ("total_inspections", "compliance_rate", "violations_by_rule",
                "trend", "severity_split", "top_offending_brands", "recent"):
        assert key in stats, "missing key: " + key
    assert len(stats["trend"]) == 30


def test_dashboard_scope_differs_by_role(client, admin_headers, inspector_headers):
    assert client.get("/api/v1/dashboard/stats", headers=admin_headers).json()["scope"] == "all"
    assert client.get("/api/v1/dashboard/stats", headers=inspector_headers).json()["scope"] == "own"


# --------------------------------------------------------------------------- #
# Reports
# --------------------------------------------------------------------------- #

def test_pdf_report_is_a_real_pdf(client, inspector_headers):
    inspection_id = scan_text(client, inspector_headers, COMPLIANT_LABEL).json()["inspection_id"]
    response = client.get("/api/v1/reports/%s?format=pdf" % inspection_id, headers=inspector_headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF-")
    assert "attachment" in response.headers["content-disposition"]


def test_pdf_contains_selectable_text_not_a_screenshot(client, inspector_headers):
    """
    The previous implementation pasted an html2canvas screenshot into a PDF, so the
    text could not be selected, searched or copied into a notice. Assert we emit
    real text operators and font resources.
    """
    inspection_id = scan_text(client, inspector_headers, COMPLIANT_LABEL).json()["inspection_id"]
    pdf = client.get("/api/v1/reports/%s?format=pdf" % inspection_id,
                     headers=inspector_headers).content
    assert b"/Font" in pdf
    assert b"/Type1" in pdf or b"/TrueType" in pdf


def test_docx_report_is_a_real_docx(client, inspector_headers):
    inspection_id = scan_text(client, inspector_headers, COMPLIANT_LABEL).json()["inspection_id"]
    response = client.get("/api/v1/reports/%s?format=docx" % inspection_id, headers=inspector_headers)
    assert response.status_code == 200
    assert response.content.startswith(b"PK")  # OOXML is a zip container

    import io
    import zipfile

    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        assert "word/document.xml" in zf.namelist()
        document = zf.read("word/document.xml").decode("utf-8")
    assert "STATUTORY COMPLIANCE INSPECTION REPORT" in document


def test_invalid_report_format_is_rejected(client, inspector_headers):
    inspection_id = scan_text(client, inspector_headers, COMPLIANT_LABEL).json()["inspection_id"]
    response = client.get("/api/v1/reports/%s?format=exe" % inspection_id, headers=inspector_headers)
    assert response.status_code == 422


def test_inspector_cannot_bulk_export(client, inspector_headers):
    assert client.get("/api/v1/reports", headers=inspector_headers).status_code == 403


def test_admin_bulk_export_is_a_real_xlsx(client, admin_headers):
    response = client.get("/api/v1/reports", headers=admin_headers)
    assert response.status_code == 200
    assert response.content.startswith(b"PK")

    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Inspections", "Violations", "Summary"]
    assert wb["Inspections"].max_row > 4


def test_report_for_another_officer_is_forbidden(client, admin_headers, inspector_headers):
    inspection_id = scan_text(client, admin_headers, COMPLIANT_LABEL).json()["inspection_id"]
    response = client.get("/api/v1/reports/%s?format=pdf" % inspection_id, headers=inspector_headers)
    assert response.status_code == 403


def test_report_filename_is_header_safe(client, inspector_headers):
    """A product name with quotes or newlines must not be able to forge a header."""
    lines = ['Evil"\r\nX-Injected: yes Product', "MRP Rs. 10 (Inclusive of all taxes)"]
    inspection_id = scan_text(client, inspector_headers, lines).json()["inspection_id"]
    response = client.get("/api/v1/reports/%s?format=pdf" % inspection_id, headers=inspector_headers)
    assert response.status_code == 200
    assert "x-injected" not in {k.lower() for k in response.headers}
    disposition = response.headers["content-disposition"]
    assert "\n" not in disposition and "\r" not in disposition
