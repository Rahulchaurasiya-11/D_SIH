"""
Editable Word compliance report (python-docx).

This is the "editable format" the problem statement asks for alongside PDF: an
officer opens the notice, adjusts the wording for the specific case, adds a file
number, and issues it. A PDF cannot serve that purpose.
"""

import io
from typing import Any, Dict

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.db import repositories as repo
from app.services.report_content import build_report_content

INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
PRIMARY = RGBColor(0x1B, 0x4D, 0x89)
DANGER = RGBColor(0xB4, 0x23, 0x18)
OK = RGBColor(0x0F, 0x7B, 0x4F)

STATUS_FILL = {
    "COMPLIANT": "0F7B4F",
    "NON_COMPLIANT": "B42318",
    "PARTIALLY_COMPLIANT": "B45309",
}


def _shade(cell, hex_colour: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(shd)


def _run(paragraph, text: str, *, bold: bool = False, size: float = 9,
         colour: RGBColor = INK, italic: bool = False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = colour
    run.font.name = "Calibri"
    return run


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    _run(p, text, bold=True, size=11, colour=PRIMARY)


def _kv_table(doc: Document, pairs, widths=(Inches(1.9), Inches(4.6))):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for label, value in pairs:
        row = table.add_row().cells
        row[0].width, row[1].width = widths
        _shade(row[0], "F1F5F9")
        _run(row[0].paragraphs[0], str(label), bold=True, size=9)
        _run(row[1].paragraphs[0], str(value), size=9)
    return table


def render_docx(inspection: Dict[str, Any]) -> bytes:
    c = build_report_content(inspection)
    doc = Document()

    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.7)
    section.top_margin = section.bottom_margin = Inches(0.7)

    # Running header
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(header_p, "GOVERNMENT OF INDIA  |  DEPARTMENT OF CONSUMER AFFAIRS", bold=True, size=8, colour=PRIMARY)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(footer_p, "Legal Metrology (Packaged Commodities) Rules, 2011  -  Reference %s"
         % c["inspection_id"], size=7, colour=MUTED)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(title, "STATUTORY COMPLIANCE INSPECTION REPORT", bold=True, size=15)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(sub, "Issued under the Legal Metrology Act, 2009 and the Legal Metrology "
              "(Packaged Commodities) Rules, 2011", size=8, colour=MUTED)

    # Verdict band
    band = doc.add_table(rows=1, cols=2)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, text in enumerate([c["status_label"], "%d / 100  COMPLIANCE SCORE" % c["overall_score"]]):
        cell = band.rows[0].cells[idx]
        _shade(cell, STATUS_FILL.get(c["status"], "64748B"))
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(cell.paragraphs[0], text, bold=True, size=12, colour=RGBColor(0xFF, 0xFF, 0xFF))

    # 1. Particulars
    _heading(doc, "1. INSPECTION PARTICULARS")
    _kv_table(doc, [
        ("Inspection ID", c["inspection_id"]),
        ("Date of inspection", c["created_at"]),
        ("Product", c["product_name"]),
        ("Brand", c["brand"]),
        ("Inspecting officer", c["officer_name"]),
        ("Jurisdiction", c["jurisdiction"]),
        ("Evidence source", c["source"].title()),
        ("Label language", c["language_profile"]),
    ])

    if c["is_manually_verified"]:
        note = doc.add_paragraph()
        _run(note, "Fields verified/corrected by the inspecting officer: %s"
             % (", ".join(c["manual_fields_applied"]) or "yes"), size=8, italic=True, colour=MUTED)

    # 2. Declarations
    _heading(doc, "2. MANDATORY DECLARATIONS DETECTED")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, head in enumerate(["Declaration required", "Value found on package", "Rule"]):
        _shade(table.rows[0].cells[i], "F1F5F9")
        _run(table.rows[0].cells[i].paragraphs[0], head, bold=True, size=9)
    for label, value, rule in c["declarations"]:
        cells = table.add_row().cells
        _run(cells[0].paragraphs[0], label, size=9)
        _run(cells[1].paragraphs[0], value, size=9,
             colour=DANGER if value == "Not detected" else INK,
             bold=value == "Not detected")
        _run(cells[2].paragraphs[0], rule, size=9)

    # 3. Violations
    _heading(doc, "3. CONTRAVENTIONS IDENTIFIED (%d)" % len(c["violations"]))
    if not c["violations"]:
        p = doc.add_paragraph()
        _run(p, "No contravention was identified. All mandatory declarations verified by the "
                "system were found to be present and in the prescribed form.", size=9, colour=OK)
    else:
        for i, v in enumerate(c["violations"], 1):
            head = doc.add_paragraph()
            head.paragraph_format.space_before = Pt(10)
            _run(head, "%d. %s" % (i, v.get("rule_name", v.get("rule_id", "Violation"))), bold=True, size=10)
            _run(head, "   [%s]" % str(v.get("severity", "MEDIUM")).upper(), bold=True, size=8, colour=DANGER)
            _kv_table(doc, [
                ("Legal reference", v.get("legal_reference", "-")),
                ("Finding", v.get("description", "-")),
                ("Evidence", v.get("found_text", "-")),
                ("Required action", v.get("remediation", "-")),
                ("Penalty band", v.get("penalty_band", "-")),
            ])

    # 4. Verified checks
    if c["passed_checks"]:
        _heading(doc, "4. DECLARATIONS VERIFIED AS COMPLIANT (%d)" % len(c["passed_checks"]))
        ok_table = doc.add_table(rows=1, cols=2)
        ok_table.style = "Table Grid"
        for i, head in enumerate(["Rule", "Verification"]):
            _shade(ok_table.rows[0].cells[i], "F1F5F9")
            _run(ok_table.rows[0].cells[i].paragraphs[0], head, bold=True, size=9)
        for chk in c["passed_checks"]:
            cells = ok_table.add_row().cells
            _run(cells[0].paragraphs[0], str(chk.get("rule_name", chk.get("rule_id", "-"))), size=9)
            _run(cells[1].paragraphs[0], str(chk.get("evidence", chk.get("description", "-"))), size=9)

    # 5. Evidence
    embedded = 0
    for item in c["evidence"][:4]:
        data = repo.evidence.load(item.get("evidence_id", ""))
        if not data:
            continue
        if embedded == 0:
            _heading(doc, "5. PHOTOGRAPHIC EVIDENCE")
        try:
            doc.add_picture(io.BytesIO(data), width=Inches(3.2))
            cap = doc.paragraphs[-1]
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(caption, "%s - %s" % (item.get("angle", "Angle"), item.get("filename", "")),
                 size=7.5, colour=MUTED, italic=True)
            embedded += 1
        except Exception:
            continue

    # Declaration + signatures
    _heading(doc, "DECLARATION")
    disc = doc.add_paragraph()
    disc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _run(disc, c["disclaimer"], size=7.5, colour=MUTED)

    doc.add_paragraph()
    sign = doc.add_table(rows=1, cols=2)
    sign.alignment = WD_TABLE_ALIGNMENT.LEFT
    for idx, (name, role) in enumerate([
        (c["officer_name"], "Inspecting Officer, Legal Metrology"),
        ("", "Controller / Authorised Officer (Countersignature)"),
    ]):
        cell = sign.rows[0].cells[idx]
        _run(cell.paragraphs[0], "\n_______________________________", size=9)
        p = cell.add_paragraph()
        _run(p, name or " ", bold=True, size=9)
        p2 = cell.add_paragraph()
        _run(p2, role, size=7.5, colour=MUTED)

    tail = doc.add_paragraph()
    _run(tail, "Report generated on %s  |  Reference: %s" % (c["generated_at"], c["inspection_id"]),
         size=7.5, colour=MUTED)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
